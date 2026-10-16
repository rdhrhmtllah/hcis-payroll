# -*- coding: utf-8 -*-
"""
Merakit "IT_BA_UAT SOP.docx" dari isi di scripts/uat_sop_data.py.

CARA KERJANYA: dokumen UAT WebKarir dibuka sebagai CETAKAN, bukan ditulis ulang
dari nol. Body-nya dikosongkan, lalu diisi kembali dengan mengkloning elemen
paragraf dan tabel aslinya.

KENAPA MENGKLONING, bukan membangun tabel baru. Ketujuh tabel pada cetakan
tidak memakai style bernama — seluruh garis, arsiran kepala tabel, lebar kolom,
dan jarak selnya ditulis langsung sebagai properti XML per tabel dan per sel.
Membangun ulang tabel baru berarti menebak semua nilai itu satu per satu sampai
hasilnya "mirip", dan mirip tidak pernah cukup untuk dokumen yang ditandatangani
bersebelahan dengan dokumen lain dari seri yang sama. Mengkloning elemennya
membuat pertanyaan itu tidak pernah perlu dijawab.

Berkas cetakan HANYA DIBACA. Ia tidak terlacak git dan tidak punya cadangan;
menimpanya akan menghapus satu-satunya salinan yang ada.

Pemakaian:
    python scripts/build_uat_sop.py
    python scripts/build_uat_sop.py --keluaran "D:/lain/UAT SOP.docx"
"""

import argparse
import copy
import os
import sys

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import uat_sop_data as isi  # noqa: E402


AKAR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CETAKAN = os.path.join(AKAR, 'IT_BA_UAT WEBKARIR-MT.docx')
KELUARAN = os.path.join(AKAR, 'IT_BA_UAT SOP.docx')

# Berapa paragraf kosong yang mendorong baris CONFIDENTIAL ke dasar halaman
# sampul.
#
# Cetakan memakai 21. Angka itu TIDAK BISA dipakai apa adanya: tabel info versi
# SOP lebih tinggi karena isinya lebih panjang, sehingga 21 spasi mendorong
# baris CONFIDENTIAL tumpah ke halaman berikutnya dan meninggalkan celah kosong
# besar di atas bab 1. Angka di bawah diukur dengan --sapu, bukan dikira-kira.
#
# Sekalipun angka ini suatu saat meleset lagi, bab 1 dipaksa mulai di halaman
# baru (lihat _paksa_halaman_baru), jadi kesalahannya paling jauh menyisakan
# ruang kosong di sampul — tidak pernah merusak susunan bab sesudahnya.
SPASI_SAMPUL = 15


# ============================================================ utilitas XML

def _teks_ke_paragraf(p_el, teks):
    """Isi ulang teks sebuah paragraf tanpa mengubah formatnya.

    Run pertama dipertahankan berikut seluruh rPr-nya (huruf, ukuran, tebal),
    run sisanya dibuang. Cetakan menyimpan formatnya di run pertama, jadi
    menulis ulang isinya di situ mempertahankan rupa aslinya.
    """
    runs = p_el.findall(qn('w:r'))

    if not runs:
        if not teks:
            return                      # sel kosong tetap kosong — tak perlu run
        r = p_el.makeelement(qn('w:r'), {})
        # Format run bawaan paragraf disimpan di w:pPr/w:rPr; bawa serta supaya
        # teks baru tidak jatuh ke huruf bawaan dokumen.
        pPr = p_el.find(qn('w:pPr'))
        if pPr is not None:
            rPr = pPr.find(qn('w:rPr'))
            if rPr is not None:
                r.append(copy.deepcopy(rPr))
        p_el.append(r)
        runs = [r]

    utama, sisa = runs[0], runs[1:]
    for r in sisa:
        p_el.remove(r)

    for anak in list(utama):
        if anak.tag != qn('w:rPr'):
            utama.remove(anak)

    for i, potong in enumerate(str(teks).split('\n')):
        if i:
            utama.append(utama.makeelement(qn('w:br'), {}))
        t = utama.makeelement(qn('w:t'), {})
        t.text = potong
        t.set(qn('xml:space'), 'preserve')
        utama.append(t)


def _teks_ke_sel(tc_el, teks):
    """Isi sebuah sel; paragraf pertama dipakai, paragraf sisanya dibuang."""
    paragraf = tc_el.findall(qn('w:p'))
    for p in paragraf[1:]:
        tc_el.remove(p)
    if paragraf:
        _teks_ke_paragraf(paragraf[0], teks)


def _teks_ke_baris(tr_el, nilai):
    sel = tr_el.findall(qn('w:tc'))
    if len(sel) != len(nilai):
        raise ValueError(
            f'Jumlah sel ({len(sel)}) tidak sama dengan jumlah nilai ({len(nilai)}). '
            'Prototipe baris yang dipilih punya bentuk berbeda dari data.'
        )
    for tc, teks in zip(sel, nilai):
        _teks_ke_sel(tc, teks)


def _ciutkan_baris(tr_el, lebar):
    """Buang sel paling kanan sebuah baris sampai tersisa len(lebar), atur lebarnya."""
    sel = tr_el.findall(qn('w:tc'))
    for tc in sel[len(lebar):]:
        tr_el.remove(tc)
    for tc, w in zip(tr_el.findall(qn('w:tc')), lebar):
        tcW = tc.find(qn('w:tcPr') + '/' + qn('w:tcW'))
        if tcW is not None:
            tcW.set(qn('w:w'), str(w))
            tcW.set(qn('w:type'), 'dxa')


def _ciutkan_kolom(tbl_el, lebar):
    """Kurangi jumlah kolom tabel jadi len(lebar), lalu pasang lebar barunya.

    Sel yang dibuang adalah sel PALING KANAN, karena hanya di situ pembuangan
    tidak menggeser makna kolom yang tersisa.

    Lebar harus dipasang di DUA tempat: w:tblGrid menentukan kerangka tabel,
    sedangkan w:tcW tiap sel menentukan lebar sel itu sendiri. Mengubah salah
    satunya saja menghasilkan tabel yang lebarnya benar tetapi kolomnya tidak
    sejajar dengan kerangkanya — dan Word menyelesaikan perselisihan itu dengan
    caranya sendiri, yang tidak selalu sama dengan yang terlihat di sini.
    """
    grid = tbl_el.find(qn('w:tblGrid'))
    kolom = grid.findall(qn('w:gridCol'))
    for gc in kolom[len(lebar):]:
        grid.remove(gc)
    for gc, w in zip(grid.findall(qn('w:gridCol')), lebar):
        gc.set(qn('w:w'), str(w))

    for tr in tbl_el.findall(qn('w:tr')):
        _ciutkan_baris(tr, lebar)


class Cetakan:
    """Penyimpan prototipe elemen yang dipanen dari dokumen cetakan."""

    def __init__(self, dok):
        p = dok.paragraphs
        self.paragraf = {
            'organisasi': p[0]._p,      # EVO GROUP — tebal 14pt, rata tengah
            'unit': p[1]._p,
            'judul': p[2]._p,           # style Title
            'subjudul': p[3]._p,        # style Subtitle
            'spasi': p[7]._p,           # paragraf kosong sampul
            'rahasia': p[28]._p,        # CONFIDENTIAL — 9pt
            'h1': p[29]._p,
            'h2': p[64]._p,
            'isi': p[30]._p,
            'butir_lebar': p[32]._p,    # butir bab 1 (indentasi lebar)
            'butir': p[56]._p,          # butir bab 5
            'butir_rapat': p[65]._p,    # butir definisi severity
            'titik': p[76]._p,          # baris titik-titik narasi

            # Pemutus section. Word menyimpan properti section DI DALAM pPr
            # paragraf terakhir section itu, bukan sebagai elemen tersendiri —
            # sehingga mengosongkan body ikut membuang seluruh pemutusnya dan
            # dokumen runtuh jadi satu section potret. Cetakan punya delapan
            # section, dan yang menampung tabel skenario adalah section
            # LANSKAP: ruang cetaknya 10,5 inci, sementara tabel sembilan
            # kolom itu selebar 10,16 inci. Tanpa pemutus ini tabelnya
            # terpotong di potret (7,5 inci) tepat pada tiga kolom yang harus
            # diisi penguji.
            'putus_potret': p[38]._p,
            'putus_lanskap': p[53]._p,
        }
        self.tabel = {
            'info': dok.tables[0]._tbl,      # 2 kolom, tanpa baris kepala
            'tiga': dok.tables[1]._tbl,      # 3 kolom, berbaris kepala
            'peserta': dok.tables[2]._tbl,   # 4 kolom
            'skenario': dok.tables[3]._tbl,  # 9 kolom
            'defect': dok.tables[4]._tbl,    # 7 kolom
            'legenda': dok.tables[5]._tbl,   # 2 kolom, tanpa baris kepala
            'signoff': dok.tables[6]._tbl,   # 6 kolom
        }

    def p(self, jenis, teks=''):
        el = copy.deepcopy(self.paragraf[jenis])
        _teks_ke_paragraf(el, teks)
        return el

    def t(self, jenis, baris, kepala=None, lebar=None):
        """Kloning tabel: baris kepala (opsional) + satu baris data per record.

        `lebar` menciutkan tabel jadi len(lebar) kolom. Penciutan harus terjadi
        DI SINI, sebelum sel diisi — bukan setelah tabelnya jadi. _teks_ke_baris
        menuntut jumlah nilai sama persis dengan jumlah sel, jadi menciutkan
        belakangan berarti mengisi sembilan sel lebih dulu dengan tujuh nilai,
        dan itu gagal sebelum sempat diciutkan.
        """
        tbl = copy.deepcopy(self.tabel[jenis])
        semua = tbl.findall(qn('w:tr'))

        proto_kepala = semua[0] if kepala is not None else None
        proto_data = semua[1] if kepala is not None else semua[0]
        proto_kepala = copy.deepcopy(proto_kepala) if proto_kepala is not None else None
        proto_data = copy.deepcopy(proto_data)

        for tr in semua:
            tbl.remove(tr)

        if lebar is not None:
            _ciutkan_kolom(tbl, lebar)          # kerangka tblGrid
            for proto in (proto_kepala, proto_data):
                if proto is not None:
                    _ciutkan_baris(proto, lebar)

        if proto_kepala is not None:
            _teks_ke_baris(proto_kepala, kepala)
            tbl.append(proto_kepala)

        for nilai in baris:
            tr = copy.deepcopy(proto_data)
            _teks_ke_baris(tr, nilai)
            tbl.append(tr)

        return tbl


def _paksa_halaman_baru(p_el):
    """Paksa paragraf ini mulai di halaman baru.

    Memakai w:pageBreakBefore, bukan menyisipkan paragraf berisi pemutus
    halaman: paragraf tambahan akan tertinggal sebagai baris kosong di puncak
    halaman berikutnya kalau susunan di atasnya berubah, sedangkan atribut ini
    menempel pada paragrafnya sendiri dan ikut ke mana pun ia bergeser.
    """
    pPr = p_el.find(qn('w:pPr'))
    if pPr is None:
        pPr = p_el.makeelement(qn('w:pPr'), {})
        p_el.insert(0, pPr)
    if pPr.find(qn('w:pageBreakBefore')) is None:
        # Urutan anak w:pPr diatur skema; pageBreakBefore harus berada di awal,
        # sesudah pStyle bila ada. Word menolak berkas dengan urutan yang salah.
        sisip = 1 if pPr.find(qn('w:pStyle')) is not None else 0
        pPr.insert(sisip, pPr.makeelement(qn('w:pageBreakBefore'), {}))
    return p_el


def _kosongkan_body(dok):
    """Buang seluruh isi body, sisakan properti section di ujungnya.

    sectPr memegang ukuran kertas, orientasi, dan margin. Ikut terbuang berarti
    dokumen jatuh ke bawaan Word, bukan ke ukuran cetakan.
    """
    body = dok.element.body
    for el in list(body.iterchildren()):
        if el.tag != qn('w:sectPr'):
            body.remove(el)


# ============================================================ perakitan


# Lebar kolom tabel uji, dalam twip. Jumlahnya harus tetap 14625 (10,16 inci) —
# itu lebar yang muat di ruang cetak section lanskap. Kolom yang dibuang dari
# bentuk sembilan-kolom: Tujuan, Prasyarat (keduanya melebur jadi poin), dan
# Hasil Aktual (Catatan sudah menampung apa yang terjadi).
#
# Kolom pertama 1550 twip (1,08 inci), bukan sekadar cukup untuk sebuah angka:
# isinya "UAT-SOP-001" yang pada 10pt menuntut ±0,9 inci ditambah margin sel.
# Kekurangan sedikit saja membuat Word mematahkannya di tanda hubung, sehingga
# tiap baris tabel bertambah tinggi satu baris teks tanpa alasan.
LEBAR_UJI = (1550, 2100, 1200, 4400, 3100, 800, 1475)

KEPALA_UJI = ('Test Case', 'Fitur', 'Penguji', 'Poin yang Diuji',
              'Expected Result', 'Status', 'Catatan')


def _baris_kasus(kasus):
    """Satu baris per fitur; poinnya jadi baris-baris di dalam satu sel.

    Poin [TOLAK] TIDAK dipisah jadi barisnya sendiri, tapi juga tidak dilebur
    jadi anak kalimat: ia berdiri sebagai poin utuh dengan penanda di depannya,
    supaya tetap terbaca sebagai hal yang harus dicoba. Sebuah fitur berstatus
    Pass hanya bila seluruh poin pada barisnya terpenuhi — aturan itu ditulis
    di KASUS_CATATAN, karena tabelnya sendiri tak bisa memaksakannya.
    """
    return [
        (kode, fitur, penguji,
         '\n'.join(f'• {t}' for t in poin),
         harapan, '', '')
        for kode, fitur, penguji, poin, harapan, _cakupan in kasus
    ]


def rakit(jalur_cetakan, jalur_keluaran, spasi_sampul=SPASI_SAMPUL):
    dok = Document(jalur_cetakan)
    c = Cetakan(dok)

    sectPr = dok.element.body.find(qn('w:sectPr'))
    _kosongkan_body(dok)
    body = dok.element.body

    def tulis(el):
        body.insert(list(body).index(sectPr), el)

    # ---------------------------------------------------------- sampul
    tulis(c.p('organisasi', isi.SAMPUL['organisasi']))
    tulis(c.p('unit', isi.SAMPUL['unit']))
    tulis(c.p('judul', isi.SAMPUL['judul']))
    tulis(c.p('subjudul', isi.SAMPUL['subjudul_1']))
    tulis(c.p('subjudul', isi.SAMPUL['subjudul_2']))
    tulis(c.p('spasi'))
    tulis(c.p('spasi'))
    tulis(c.t('info', isi.INFO_DOKUMEN))
    for _ in range(spasi_sampul):
        tulis(c.p('spasi'))
    tulis(c.p('rahasia', isi.SAMPUL['rahasia']))

    # ---------------------------------------------------------- 1. Tujuan
    tulis(_paksa_halaman_baru(c.p('h1', '1. Tujuan UAT')))
    tulis(c.p('isi', isi.TUJUAN_PEMBUKA))
    tulis(c.p('isi', isi.TUJUAN_PENGANTAR))
    for butir in isi.TUJUAN_BUTIR:
        tulis(c.p('butir_lebar', butir))
    tulis(c.p('isi'))

    # ---------------------------------------------------------- 2. Ruang Lingkup
    tulis(c.p('h1', '2. Ruang Lingkup UAT'))
    tulis(c.t(
        'tiga',
        [(str(i), nama, fokus) for i, (nama, fokus) in enumerate(isi.RUANG_LINGKUP, 1)],
        kepala=('No', 'Modul', 'Fokus Validasi Bisnis'),
    ))
    tulis(c.p('isi'))
    tulis(c.p('isi', isi.DILUAR_LINGKUP))

    # ---------------------------------------------------------- 3. Peserta
    tulis(c.p('h1', '3. Peserta UAT'))
    tulis(c.p('isi', isi.PESERTA_PEMBUKA))
    tulis(c.t(
        'peserta',
        [(str(i), nama, jabatan, peran)
         for i, (nama, jabatan, peran) in enumerate(isi.PESERTA, 1)],
        kepala=('No', 'Nama', 'Jabatan', 'Role dalam UAT'),
    ))
    tulis(c.p('isi'))

    # ---------------------------------------------------------- 3a. Prasyarat
    tulis(c.p('h2', 'Prasyarat & Data Uji'))
    tulis(c.p('isi', isi.PRASYARAT_PEMBUKA))
    tulis(c.t(
        'tiga',
        [(str(i), nama, ket) for i, (nama, ket) in enumerate(isi.PRASYARAT, 1)],
        kepala=('No', 'Prasyarat', 'Keterangan'),
    ))
    # Tutup bagian potret sebelum bab 4 berpindah ke lanskap.
    tulis(c.p('putus_potret'))

    # ------------------------------------------------------- 4. Daftar uji
    # SELURUH bab ini berada di section lanskap; pemutusnya dipasang setelah
    # tabel, karena sectPr menggambarkan section yang BERAKHIR pada paragraf
    # tempat ia menempel — bukan yang dimulai.
    #
    # Daftar 48 skenario rinci TIDAK dicetak. Ia sempat ikut sebagai lampiran,
    # dan hasilnya dokumen justru bertambah tebal padahal permintaannya
    # menyederhanakan: 14 baris fitur ditambah 48 baris rincian. Rinciannya
    # tetap hidup di isi.SKENARIO dan di riwayat git bila suatu saat perlu
    # ditelusuri.
    tulis(c.p('h1', '4. Daftar Uji SOP'))
    tulis(c.p('isi', isi.KASUS_PEMBUKA))
    tulis(c.p('isi', isi.KASUS_CATATAN))
    tulis(c.t('skenario', _baris_kasus(isi.KASUS),
              kepala=KEPALA_UJI, lebar=LEBAR_UJI))
    tulis(c.p('putus_lanskap'))

    # ---------------------------------------------------------- 5. Kriteria
    tulis(c.p('h1', '5. UAT Acceptance Criteria'))
    tulis(c.p('isi', isi.KRITERIA_PEMBUKA))
    for butir in isi.KRITERIA:
        tulis(c.p('butir', butir))

    # ---------------------------------------------------------- 6. Defect Log
    tulis(c.p('h1', '6. Defect / Issue Log'))
    tulis(c.p('isi', isi.DEFECT_PEMBUKA))
    tulis(c.t(
        'defect',
        [('', '', '', '', '', '', '')] * isi.DEFECT_BARIS_KOSONG,
        kepala=('No', 'UAT-ID', 'Deskripsi Issue', 'Severity', 'PIC', 'Status',
                'Tanggal'),
    ))
    tulis(c.p('h2', 'Definisi Severity:'))
    for butir in isi.SEVERITY:
        tulis(c.p('butir_rapat', butir))
    tulis(c.p('isi'))

    # ---------------------------------------------------------- 7. Hasil
    tulis(c.p('h1', '7. Hasil & Kesimpulan UAT'))
    tulis(c.t('legenda', isi.HASIL_KESIMPULAN))
    tulis(c.p('isi'))
    tulis(c.p('h2', 'Narasi Kesimpulan'))
    tulis(c.p('isi', isi.NARASI_PEMBUKA))
    for _ in range(isi.NARASI_BARIS):
        tulis(c.p('titik', isi.NARASI_TITIK))
    tulis(c.p('isi'))

    # ---------------------------------------------------------- 8. Sign-Off
    tulis(c.p('h1', '8. Sign-Off UAT'))
    tulis(c.p('isi', isi.SIGNOFF_PEMBUKA))
    tulis(c.t(
        'signoff',
        [(str(i), nama, jabatan, peran, '', '')
         for i, (nama, jabatan, peran) in enumerate(isi.PESERTA, 1)],
        kepala=('No', 'Nama', 'Jabatan', 'Role dalam UAT', 'Tanda Tangan', 'Tanggal'),
    ))

    dok.save(jalur_keluaran)
    return jalur_keluaran


# Daftar berkas dilewatkan lewat berkas daftar, bukan argumen: powershell
# -Command tidak menerima -args (itu hanya untuk -File), dan menempelkan jalur
# langsung ke dalam string perintah membuat jalur berspasi terpecah.
_PS_UKUR = r'''
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false; $w.DisplayAlerts = 0
try {
  foreach ($f in (Get-Content -LiteralPath $env:UAT_DAFTAR)) {
    $d = $w.Documents.Open($f, [ref]$false, [ref]$true, [ref]$false)
    # Find dijalankan pada Range dan MENIMPA Range itu dengan hasilnya. Membaca
    # Selection di sini keliru: Selection tidak ikut bergerak, sehingga tiap
    # dokumen akan melapor "halaman 1" apa pun keadaannya.
    $r = $d.Content
    $r.Find.Text = 'CONFIDENTIAL'
    $ketemu = $r.Find.Execute()
    # wdActiveEndPageNumber = 3
    $hal = if ($ketemu) { $r.Information(3) } else { -1 }
    "$([System.IO.Path]::GetFileNameWithoutExtension($f))=$hal/$($d.ComputeStatistics(2))"
    $d.Close([ref]$false)
  }
} finally { $w.Quit() }
'''


def sapu(jalur_cetakan, rentang):
    """Cari jumlah spasi sampul terbesar yang masih menyisakan CONFIDENTIAL di
    halaman 1.

    Diukur dengan merender lewat Word, bukan dihitung dari tinggi baris:
    tinggi sel bergantung pada pembungkusan teks, dan pembungkusan bergantung
    pada font yang benar-benar dipakai saat merender.
    """
    import subprocess
    import tempfile

    tmp = tempfile.mkdtemp(prefix='uat-sop-sapu-')
    berkas = []
    for n in rentang:
        jalur = os.path.join(tmp, f'{n:02d}.docx')
        rakit(jalur_cetakan, jalur, spasi_sampul=n)
        berkas.append(jalur)

    daftar = os.path.join(tmp, 'daftar.txt')
    with open(daftar, 'w', encoding='utf-8') as f:
        f.write('\n'.join(berkas))

    keluar = subprocess.run(
        ['powershell', '-NoProfile', '-Command', _PS_UKUR],
        capture_output=True, text=True,
        env={**os.environ, 'UAT_DAFTAR': daftar},
    )
    if keluar.returncode != 0:
        raise SystemExit('Gagal mengukur lewat Word:\n' + keluar.stderr)

    print(f'{"spasi":>6}  {"hal CONFIDENTIAL":>18}  {"total hal":>9}')
    terbaik = None
    for baris in keluar.stdout.strip().splitlines():
        if '=' not in baris:
            continue
        n, angka = baris.split('=')
        hal, total = angka.split('/')
        tanda = ''
        if int(hal) == 1:
            terbaik = max(terbaik or 0, int(n))
            tanda = '  <- muat di sampul'
        print(f'{int(n):>6}  {int(hal):>18}  {int(total):>9}{tanda}')

    print()
    if terbaik is None:
        print('Tidak ada yang muat. Kurangi rentangnya.')
    else:
        print(f'Setel SPASI_SAMPUL = {terbaik}')
    return terbaik


def main():
    ap = argparse.ArgumentParser(description='Rakit dokumen UAT Modul SOP.')
    ap.add_argument('--cetakan', default=CETAKAN)
    ap.add_argument('--keluaran', default=KELUARAN)
    ap.add_argument('--spasi', type=int, default=SPASI_SAMPUL,
                    help='jumlah paragraf kosong pada sampul')
    ap.add_argument('--sapu', action='store_true',
                    help='ukur jumlah spasi sampul terbaik lewat Word, '
                         'lalu berhenti tanpa menulis keluaran')
    arg = ap.parse_args()

    if not os.path.exists(arg.cetakan):
        raise SystemExit(f'Cetakan tidak ditemukan: {arg.cetakan}')

    if arg.sapu:
        sapu(arg.cetakan, range(10, 23))
        return

    if os.path.abspath(arg.cetakan) == os.path.abspath(arg.keluaran):
        raise SystemExit(
            'Keluaran menunjuk berkas cetakan itu sendiri. Cetakan tidak terlacak '
            'git dan tidak punya cadangan — menimpanya menghapus satu-satunya '
            'salinan yang ada.'
        )

    hasil = rakit(arg.cetakan, arg.keluaran, spasi_sampul=arg.spasi)
    print(f'Tersimpan: {hasil}')
    poin = sum(len(k[3]) for k in isi.KASUS)
    tolak = sum(1 for k in isi.KASUS for t in k[3] if t.startswith('[TOLAK]'))
    # Yang dilaporkan adalah yang DICETAK. Melaporkan 48 skenario padahal yang
    # masuk dokumen 14 baris fitur pernah menyembunyikan persis kesalahan yang
    # membuat dokumen ini membengkak.
    print(f'  {len(isi.RUANG_LINGKUP)} modul, {len(isi.KASUS)} fitur diuji, '
          f'{poin} poin ({tolak} di antaranya poin [TOLAK])')


if __name__ == '__main__':
    main()
